import os
import sys
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def collect_files_with_hierarchy(root_dir):
    """
    遍历目录，收集每个文件的层级文件夹列表和文件名
    返回: (file_rows, max_depth)
        file_rows: 列表，每个元素为 (folder_list, filename)
            folder_list: 从根下第一级开始的各级文件夹名列表（按顺序）
            filename: 文件名（含后缀）
        max_depth: 最大文件夹层级数（即folder_list的最大长度）
    """
    root_dir = os.path.abspath(root_dir)
    if not os.path.isdir(root_dir):
        print(f"错误：'{root_dir}' 不是有效的目录。")
        return None, 0

    file_rows = []
    max_depth = 0
    for current_dir, _, files in os.walk(root_dir):
        for file in files:
            full_path = os.path.join(current_dir, file)
            rel_path = os.path.relpath(full_path, root_dir)
            # 拆分为文件夹部分和文件名
            parts = rel_path.split(os.sep)
            # 最后一个是文件名，前面的都是文件夹
            *folders, filename = parts
            file_rows.append((folders, filename))
            max_depth = max(max_depth, len(folders))
    return file_rows, max_depth

def write_hierarchy_to_excel(file_rows, max_depth, root_name, output_path):
    """
    将层级数据写入 Excel，格式为：根文件夹 | 子文件夹1 | 子文件夹2 | ... | 文件名
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "文件清单"

    # 生成表头
    headers = ["根文件夹"] + [f"子文件夹{i+1}" for i in range(max_depth)] + ["文件名"]
    ws.append(headers)

    # 写入数据行
    for folders, filename in file_rows:
        row = [root_name]  # 第一列：根文件夹名
        # 填充各级子文件夹（最多 max_depth 列）
        for i in range(max_depth):
            if i < len(folders):
                row.append(folders[i])
            else:
                row.append("")  # 留空
        # 最后一列：文件名
        row.append(filename)
        ws.append(row)

    # 自动调整列宽
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = max_length + 2
        ws.column_dimensions[col_letter].width = adjusted_width

    wb.save(output_path)
    print(f"成功生成文件清单：{output_path}")
    print(f"共找到 {len(file_rows)} 个文件。")

if __name__ == "__main__":
    # 获取目标文件夹路径（命令行参数，默认为当前目录）
    if len(sys.argv) > 1:
        target_dir = sys.argv[1]
    else:
        target_dir = os.getcwd()

    # 定义输出 Excel 文件名（固定为“XK060文件清单.xlsx”）
    output_file = os.path.join(target_dir, "XK060文件清单.xlsx")

    # 收集文件层级数据
    file_rows, max_depth = collect_files_with_hierarchy(target_dir)
    if file_rows is not None:
        # 根文件夹名称（目标文件夹的最后一节）
        root_name = os.path.basename(os.path.abspath(target_dir))
        # 如果根文件夹名为空（如驱动器根目录），则使用盘符
        if not root_name:
            root_name = target_dir  # 如 "C:" 或 "/"
        # 按相对路径排序（可选，使输出更有序）
        file_rows.sort(key=lambda x: os.path.join(*x[0], x[1]))
        write_hierarchy_to_excel(file_rows, max_depth, root_name, output_file)