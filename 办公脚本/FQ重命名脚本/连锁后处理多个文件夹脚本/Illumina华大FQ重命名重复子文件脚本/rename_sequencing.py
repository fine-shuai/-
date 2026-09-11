#!/usr/bin/env python3
"""
二代测序数据文件批量重命名脚本（增强版）
支持华大MGI T7和Illumina两种数据格式
支持根据子文件夹名称区分重复的原文件名（通过mapping表的“文件夹名称”列）
"""

import os
import sys
import pandas as pd
import argparse
import re
import shutil
from pathlib import Path

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("警告: 未安装 openpyxl，Excel 格式保留功能可能受限。建议安装: pip install openpyxl")


# ==================== 公共函数 ====================
def parse_arguments():
    parser = argparse.ArgumentParser(description='二代测序数据文件批量重命名脚本（增强版）')
    parser.add_argument('-m', '--mapping-file', required=True,
                        help='映射表格文件路径 (Excel格式)')
    parser.add_argument('-d', '--data-dir', required=True,
                        help='包含测序数据的根目录路径')
    parser.add_argument('--mode', choices=['mgi', 'illumina'],
                        help='处理模式: mgi(华大) 或 illumina')
    parser.add_argument('--dry-run', action='store_true',
                        help='试运行，不实际修改文件')
    parser.add_argument('--overwrite', action='store_true',
                        help='覆盖已存在的目标文件')
    parser.add_argument('--batch', action='store_true',
                        help='强制批量模式：将-d指定的目录视为根文件夹，处理其下所有子文件夹')
    parser.add_argument('-c', '--original-col', default='原文件名',
                        help='映射表中原文件名列名 (默认: 原文件名)')
    parser.add_argument('-n', '--new-col', default='新文件名',
                        help='映射表中新文件名列名 (默认: 新文件名)')
    parser.add_argument('-f', '--folder-col', default='文件夹名称',
                        help='映射表中文件夹名称列名 (默认: 文件夹名称)')
    parser.add_argument('--sample-barcode-col', default='样本条码',
                        help='华大模式中Excel表格样本条码列名 (默认: 样本条码)')
    return parser.parse_args()


def read_mapping_file(file_path, folder_col, original_col, new_col):
    """读取映射表格文件，返回两个映射字典"""
    try:
        df = pd.read_excel(file_path)
        required_cols = [original_col, new_col]
        for col in required_cols:
            if col not in df.columns:
                raise ValueError(f"映射表中找不到列: {col}")

        folder_mappings = {}
        global_mappings = {}

        for _, row in df.iterrows():
            original_name = str(row[original_col]).strip()
            new_name = str(row[new_col]).strip()
            if not original_name or original_name == 'nan' or not new_name or new_name == 'nan':
                continue

            # 处理文件夹名称列
            folder_val = row.get(folder_col)
            if pd.isna(folder_val) or str(folder_val).strip() == '':
                global_mappings[original_name] = new_name
            else:
                folder_str = str(folder_val).strip()
                # 处理数字格式（去除 .0）
                if '.' in folder_str and folder_str.replace('.', '', 1).isdigit():
                    folder_str = str(int(float(folder_str)))
                if folder_str not in folder_mappings:
                    folder_mappings[folder_str] = {}
                folder_mappings[folder_str][original_name] = new_name

        return folder_mappings, global_mappings
    except Exception as e:
        print(f"读取映射文件时出错: {e}")
        return {}, {}


def get_mapping_for_folder(folder_name, folder_mappings, global_mappings):
    """获取当前文件夹适用的映射字典（优先文件夹特定映射，然后全局映射）"""
    mapping = {}
    mapping.update(global_mappings)
    if folder_name in folder_mappings:
        mapping.update(folder_mappings[folder_name])
    return mapping


def rename_fq_files_generic(data_dir, mapping, overwrite, dry_run, is_illumina):
    """
    通用文件重命名函数
    is_illumina: True表示Illumina模式（按前缀匹配），False表示华大模式（按前缀匹配+去除_raw）
    """
    fq_files = []
    for ext in ['*.fastq', '*.fq', '*.fastq.gz', '*.fq.gz']:
        fq_files.extend(Path(data_dir).glob(ext))

    if not fq_files:
        print(f"在目录 {data_dir} 中未找到fq文件")
        return 0

    print(f"找到 {len(fq_files)} 个fq文件")
    renamed_count = 0

    for fq_file in fq_files:
        file_name = fq_file.name
        new_file_name = None

        if is_illumina:
            # Illumina模式：前缀匹配
            for old_name, new_name in mapping.items():
                if file_name.startswith(old_name):
                    suffix = file_name[len(old_name):]
                    new_file_name = new_name + suffix
                    break
        else:
            # 华大模式：前缀匹配，并去除_raw
            for old_name, new_name in mapping.items():
                if file_name.startswith(old_name):
                    suffix = file_name[len(old_name):]
                    new_file_name = new_name + suffix
                    # 去除 _raw 部分（华大特有）
                    new_file_name = re.sub(r'_raw(?=_[12]\.fq\.gz)', '', new_file_name)
                    break

        if not new_file_name:
            print(f"文件 {file_name} 不匹配任何映射关系，跳过")
            continue

        new_file_path = os.path.join(data_dir, new_file_name)
        print(f"原文件: {file_name}")
        print(f"新文件: {new_file_name}")

        if not dry_run:
            try:
                if os.path.exists(new_file_path):
                    if overwrite:
                        os.remove(new_file_path)
                        print(f"已删除已存在的文件: {new_file_name}")
                    else:
                        print(f"警告: 目标文件已存在，跳过重命名 (使用 --overwrite 覆盖)")
                        continue
                os.rename(fq_file, new_file_path)
                print(f"重命名成功")
                renamed_count += 1
            except Exception as e:
                print(f"重命名失败: {e}")
        else:
            renamed_count += 1

        print("-" * 50)

    return renamed_count


def update_excel_file_mgi(data_dir, mapping, sample_barcode_col, dry_run=False):
    """更新华大 MGI 目录中的 Excel 文件（.xls 或 .xlsx），使用当前文件夹的映射"""
    excel_files = list(Path(data_dir).glob("*.xls")) + list(Path(data_dir).glob("*.xlsx"))
    excel_files = [f for f in excel_files if not f.name.startswith('~$')]

    if not excel_files:
        print(f"在目录 {data_dir} 中未找到Excel文件")
        return 0

    if len(excel_files) > 1:
        print(f"警告: 找到多个Excel文件，将处理第一个: {excel_files[0].name}")

    excel_file = excel_files[0]
    file_name = excel_file.name
    print(f"处理Excel文件: {file_name}")

    try:
        if HAS_OPENPYXL and excel_file.suffix == '.xlsx':
            wb = load_workbook(filename=excel_file)
            if 'Sheet1' not in wb.sheetnames:
                print(f"文件 {file_name} 中没有 'Sheet1' 工作表，跳过")
                return 0

            ws = wb['Sheet1']
            barcode_col = None
            for col in range(1, ws.max_column + 1):
                if ws.cell(row=1, column=col).value == sample_barcode_col:
                    barcode_col = col
                    break

            if barcode_col is None:
                print(f"文件 {file_name} 的Sheet1工作表中没有 '{sample_barcode_col}' 列，跳过")
                return 0

            changes_made = False
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=barcode_col)
                if cell.value is None:
                    continue
                value_str = str(cell.value).strip()
                for old_name, new_name in mapping.items():
                    if value_str == old_name or old_name in value_str:
                        new_value = value_str.replace(old_name, new_name)
                        new_value = re.sub(r'_raw$', '', new_value)
                        cell.value = new_value
                        print(f"更新样本条码: {value_str} -> {new_value}")
                        changes_made = True
                        break

            if changes_made:
                if not dry_run:
                    backup_file = excel_file.with_suffix(excel_file.suffix + '.bak')
                    shutil.copy2(excel_file, backup_file)
                    print(f"已创建备份文件: {backup_file.name}")
                    wb.save(excel_file)
                    print(f"Excel文件更新成功: {file_name}")
                    return 1
                else:
                    print(f"计划更新Excel文件: {file_name}")
                    return 1
            else:
                print(f"Excel文件无需更新: {file_name}")
                return 0

        else:
            # 回退使用 pandas
            excel_data = pd.read_excel(excel_file, sheet_name=None)
            if 'Sheet1' not in excel_data:
                print(f"文件 {file_name} 中没有 'Sheet1' 工作表，跳过")
                return 0

            df = excel_data['Sheet1']
            if sample_barcode_col not in df.columns:
                print(f"文件 {file_name} 的Sheet1工作表中没有 '{sample_barcode_col}' 列，跳过")
                return 0

            changes_made = False
            for i, value in enumerate(df[sample_barcode_col]):
                if pd.isna(value):
                    continue
                value_str = str(value).strip()
                for old_name, new_name in mapping.items():
                    if value_str == old_name or old_name in value_str:
                        new_value = value_str.replace(old_name, new_name)
                        new_value = re.sub(r'_raw$', '', new_value)
                        df.at[i, sample_barcode_col] = new_value
                        print(f"更新样本条码: {value_str} -> {new_value}")
                        changes_made = True
                        break

            if changes_made:
                if not dry_run:
                    backup_file = excel_file.with_suffix(excel_file.suffix + '.bak')
                    shutil.copy2(excel_file, backup_file)
                    print(f"已创建备份文件: {backup_file.name}")

                    excel_data['Sheet1'] = df
                    with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
                        for sheet_name, sheet_data in excel_data.items():
                            sheet_data.to_excel(writer, sheet_name=sheet_name, index=False)

                    print(f"Excel文件更新成功: {file_name}")
                    return 1
                else:
                    print(f"计划更新Excel文件: {file_name}")
                    return 1
            else:
                print(f"Excel文件无需更新: {file_name}")
                return 0

    except Exception as e:
        print(f"处理Excel文件 {file_name} 时出错: {e}")
        return 0


def process_single_directory(data_dir, folder_mappings, global_mappings, mode, sample_barcode_col, overwrite, dry_run):
    """处理单个目录，根据当前文件夹名称选择映射"""
    folder_name = Path(data_dir).name
    mapping = get_mapping_for_folder(folder_name, folder_mappings, global_mappings)
    if not mapping:
        print(f"警告: 文件夹 {folder_name} 没有适用的映射关系，跳过")
        return 0, 0

    print(f"\n{'='*60}")
    print(f"处理目录: {data_dir} (文件夹名: {folder_name})")
    print(f"使用 {len(mapping)} 条映射规则")
    print('='*60)

    if mode == 'mgi':
        fq_renamed = rename_fq_files_generic(data_dir, mapping, overwrite, dry_run, is_illumina=False)
        excel_updated = update_excel_file_mgi(data_dir, mapping, sample_barcode_col, dry_run)
        if dry_run:
            print(f"试运行完成，计划重命名 {fq_renamed} 个fq文件，更新 {excel_updated} 个Excel文件")
        else:
            print(f"处理完成，重命名 {fq_renamed} 个fq文件，更新 {excel_updated} 个Excel文件")
        return fq_renamed, excel_updated
    else:  # illumina
        fq_renamed = rename_fq_files_generic(data_dir, mapping, overwrite, dry_run, is_illumina=True)
        if dry_run:
            print(f"试运行完成，计划重命名 {fq_renamed} 个fq文件")
        else:
            print(f"处理完成，重命名 {fq_renamed} 个fq文件")
        return fq_renamed, 0


def get_directories_to_process(root_dir, batch_mode):
    """确定要处理的目录列表"""
    root = Path(root_dir)
    if batch_mode:
        subdirs = [d for d in root.iterdir() if d.is_dir() and not d.name.startswith('.')]
        if not subdirs:
            print(f"在根目录 {root_dir} 中未找到任何子文件夹")
            return []
        return subdirs
    else:
        has_fq = any(root.glob('*.fastq*')) or any(root.glob('*.fq*'))
        if has_fq:
            return [root]
        else:
            subdirs = [d for d in root.iterdir() if d.is_dir() and not d.name.startswith('.')]
            if subdirs:
                print("当前目录无fq文件，自动处理所有子文件夹。")
                return subdirs
            else:
                print("当前目录无fq文件且无子文件夹，无法处理。")
                return []


def main():
    # 处理拖放参数
    if len(sys.argv) > 1 and not sys.argv[1].startswith('-'):
        data_dir = sys.argv[1]
        sys.argv = [sys.argv[0]] + sys.argv[2:]
        sys.argv.extend(['-d', data_dir])

    args = parse_arguments()

    print("二代测序数据文件批量重命名脚本（增强版）")
    print("=" * 70)
    print("功能: 根据映射表重命名测序文件，支持子文件夹名称区分重复原文件名")
    print("=" * 70)

    # 读取映射文件
    folder_mappings, global_mappings = read_mapping_file(
        args.mapping_file, args.folder_col, args.original_col, args.new_col
    )
    if not folder_mappings and not global_mappings:
        print("未找到有效的映射关系，请检查映射文件")
        return

    print(f"成功读取全局映射 {len(global_mappings)} 条，文件夹特定映射涉及 {len(folder_mappings)} 个文件夹")

    # 确定模式
    mode = args.mode
    if mode is None:
        print("\n请选择测序平台模式:")
        print("1. 华大 MGI (mgi)")
        print("2. Illumina")
        choice = input("请输入数字 (1 或 2): ").strip()
        if choice == '1':
            mode = 'mgi'
        elif choice == '2':
            mode = 'illumina'
        else:
            print("输入无效，默认使用华大模式")
            mode = 'mgi'

    print(f"\n模式: {'华大 MGI' if mode == 'mgi' else 'Illumina'}")

    # 获取要处理的目录列表
    dirs_to_process = get_directories_to_process(args.data_dir, args.batch)
    if not dirs_to_process:
        return

    # 统计
    total_fq = 0
    total_excel = 0
    for sub_dir in dirs_to_process:
        fq_cnt, excel_cnt = process_single_directory(
            sub_dir, folder_mappings, global_mappings, mode,
            args.sample_barcode_col, args.overwrite, args.dry_run
        )
        total_fq += fq_cnt
        total_excel += excel_cnt

    print("\n" + "="*70)
    if args.dry_run:
        print(f"试运行总结: 共计划重命名 {total_fq} 个fq文件，更新 {total_excel} 个Excel文件")
    else:
        print(f"处理总结: 共重命名 {total_fq} 个fq文件，更新 {total_excel} 个Excel文件")
    print("="*70)

    if os.name == 'nt':
        input("\n按Enter键退出...")


if __name__ == '__main__':
    main()